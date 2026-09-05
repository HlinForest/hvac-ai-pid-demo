
from __future__ import annotations
import csv
from datetime import date
from pathlib import Path
import math
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs" / "documents"
OUT.mkdir(parents=True, exist_ok=True)
FONT = "Microsoft YaHei"
CODE = "Consolas"
BLUE, NAVY, DARK, MUTED = "2E74B5", "0B2545", "1F4D78", "667085"
LIGHT, GRAY, CALLOUT, GOLD = "E8EEF5", "F2F4F7", "F4F6F9", "7A5A00"

def rows(name):
    with (ROOT / name).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def num(x, n=5):
    try:
        x = float(x)
    except (TypeError, ValueError):
        return str(x)
    if not math.isfinite(x):
        return "—"
    if abs(x) >= 1000 or (0 < abs(x) < 0.001):
        return f"{x:.{n}g}"
    return f"{x:.{n}f}".rstrip("0").rstrip(".")

def runfmt(run, font=FONT, size=10.5, bold=False, italic=False, color="1F2937"):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), font)
    rpr.rFonts.set(qn("w:hAnsi"), font)
    rpr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic
    run.font.color.rgb = RGBColor.from_string(color)

def stylefmt(style, size=10.5, color="1F2937", bold=False, after=6, before=0, line=1.15):
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line

def margins(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar"); tcpr.append(mar)
    for side, val in (("top",80),("start",120),("bottom",80),("end",120)):
        node = mar.find(qn("w:"+side))
        if node is None:
            node = OxmlElement("w:"+side); mar.append(node)
        node.set(qn("w:w"), str(val)); node.set(qn("w:type"), "dxa")

def shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcpr.append(shd)
    shd.set(qn("w:fill"), fill)

def geometry(table, widths, indent=120):
    total = sum(widths)
    pr = table._tbl.tblPr
    tw = pr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); pr.insert(0, tw)
    tw.set(qn("w:w"), str(total)); tw.set(qn("w:type"), "dxa")
    ti = pr.find(qn("w:tblInd"))
    if ti is None:
        ti = OxmlElement("w:tblInd"); pr.append(ti)
    ti.set(qn("w:w"), str(indent)); ti.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW"); tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths[i])); tcw.set(qn("w:type"), "dxa")
            margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.LEFT

def celltext(cell, text, size=8.3, bold=False, color="1F2937", font=FONT):
    cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    runfmt(p.add_run(str(text)), font=font, size=size, bold=bold, color=color)

def table(doc, headers, data, widths=None, size=8.3):
    widths = widths or [9360//len(headers)]*len(headers)
    widths[-1] += 9360-sum(widths)
    t = doc.add_table(rows=1, cols=len(headers)); geometry(t, widths)
    trpr = t.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trpr.append(hdr)
    for c, h in zip(t.rows[0].cells, headers, strict=True):
        shading(c, LIGHT); celltext(c, h, size, True, NAVY)
    for values in data:
        cells = t.add_row().cells
        for c, v in zip(cells, values, strict=True): celltext(c, v, size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def para(doc, text="", *, bold=False, color="1F2937", italic=False, align=None, after=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    runfmt(p.add_run(text), bold=bold, color=color, italic=italic)
    return p

def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet"); runfmt(p.add_run(text), size=10.5); return p

def step(doc, text):
    p = doc.add_paragraph(style="List Number"); runfmt(p.add_run(text), size=10.5); return p

def code(doc, text, caption=None):
    if caption: para(doc, caption, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    t = doc.add_table(rows=1, cols=1); geometry(t, [9360]); c=t.cell(0,0); shading(c,"F7F8FA"); c.text=""
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    for i, line in enumerate(text.strip().splitlines()):
        if i: p.add_run().add_break()
        runfmt(p.add_run(line), font=CODE, size=8.2, color="263238")
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def _math_run(text):
    run = OxmlElement("m:r")
    rpr = OxmlElement("m:rPr")
    style = OxmlElement("m:sty")
    style.set(qn("m:val"), "p")
    rpr.append(style); run.append(rpr)
    node = OxmlElement("m:t")
    node.set(qn("xml:space"), "preserve")
    node.text = str(text)
    run.append(node)
    return run

def _math_seq(parent, pieces):
    for piece in pieces:
        if isinstance(piece, str):
            parent.append(_math_run(piece)); continue
        kind = piece[0]
        if kind in ("sub", "sup"):
            node = OxmlElement("m:sSub" if kind == "sub" else "m:sSup")
            base = OxmlElement("m:e"); sub = OxmlElement("m:sub" if kind == "sub" else "m:sup")
            _math_seq(base, [piece[1]] if isinstance(piece[1], str) else piece[1])
            _math_seq(sub, [piece[2]] if isinstance(piece[2], str) else piece[2])
            node.extend([base, sub]); parent.append(node)
        elif kind == "frac":
            node = OxmlElement("m:f")
            num = OxmlElement("m:num"); den = OxmlElement("m:den")
            _math_seq(num, piece[1]); _math_seq(den, piece[2])
            node.extend([num, den]); parent.append(node)
        elif kind == "supsub":
            node = OxmlElement("m:sSubSup")
            base = OxmlElement("m:e"); sub = OxmlElement("m:sub"); sup = OxmlElement("m:sup")
            _math_seq(base, [piece[1]]); _math_seq(sub, [piece[2]]); _math_seq(sup, [piece[3]])
            node.extend([base, sub, sup]); parent.append(node)
        elif kind == "int":
            parent.append(_math_run("∫"))
            _math_seq(parent, [piece[1]])

def equation(doc, pieces, caption=None):
    if caption: para(doc, caption, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(8)
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    _math_seq(omath, pieces)
    omath_para.append(omath); p._p.append(omath_para)

def callout(doc, label, text, fill=CALLOUT):
    t=doc.add_table(rows=1, cols=1); geometry(t,[9360])
    trpr=t.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement("w:tblHeader"); hdr.set(qn("w:val"),"true"); trpr.append(hdr)
    c=t.cell(0,0); shading(c,fill); c.text=""
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    runfmt(p.add_run(label+"："), bold=True, color=NAVY, size=10.5); runfmt(p.add_run(text), size=10.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)

def picture(doc, rel, caption, width=6.2):
    path=ROOT/rel
    if path.exists():
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=p.add_run(); run.add_picture(str(path), width=Inches(width))
        drawing=run._r.xpath('.//wp:docPr')
        if drawing:
            drawing[0].set('descr', caption); drawing[0].set('title', caption)
        para(doc, caption, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    else:
        callout(doc, "图件缺失", rel, fill="FFF4E5")

def source(doc, text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(4)
    runfmt(p.add_run("数据来源："), size=8.5, color=MUTED, bold=True); runfmt(p.add_run(text), size=8.5, color=MUTED)

def setup(doc, label, compact=False):
    sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(0.492)
    stylefmt(doc.styles["Normal"], 10.5 if compact else 11, after=6, line=1.25 if compact else 1.1)
    stylefmt(doc.styles["Title"], 25 if compact else 24, NAVY, True, after=7, line=1.05)
    stylefmt(doc.styles["Subtitle"], 13, MUTED, after=14)
    stylefmt(doc.styles["Heading 1"], 16, BLUE, True, before=18 if compact else 16, after=8)
    stylefmt(doc.styles["Heading 2"], 13, BLUE, True, before=14 if compact else 12, after=6)
    stylefmt(doc.styles["Heading 3"], 11.5, DARK, True, before=10, after=4)
    for s in ("List Bullet","List Number"):
        stylefmt(doc.styles[s], 10.5, after=4, line=1.2 if compact else 1.167)
    hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; hp.text=""; runfmt(hp.add_run(label), size=8.5, color=MUTED)
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; fp.text=""; runfmt(fp.add_run("页码 "), size=9, color=MUTED)
    r=fp.add_run(); fld=OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"),"begin"); ins=OxmlElement("w:instrText"); ins.set(qn("xml:space"),"preserve"); ins.text="PAGE"; end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end"); r._r.extend([fld,ins,end]); runfmt(r,size=9,color=MUTED)

def title(doc, name, sub, meta, beginner=False):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(28 if beginner else 18); p.paragraph_format.space_after=Pt(6); runfmt(p.add_run("零基础解释版" if beginner else "提交报告"), size=10.5, color=GOLD if beginner else BLUE, bold=True)
    p=doc.add_paragraph(style="Title"); runfmt(p.add_run(name), size=25 if beginner else 24, color=NAVY, bold=True)
    p=doc.add_paragraph(style="Subtitle"); runfmt(p.add_run(sub), size=13, color=MUTED)
    data=[[k,v] for k,v in meta]; table(doc,["项目","说明"],data,[1800,7560],9.1)

def common(doc, beginner=False):
    heading(doc,"1. 先建立共同语言：这些算法到底在控制什么？")
    para(doc,"把房间想成一个会慢慢升温、也会慢慢降温的大盒子。空调不会瞬间改变温度；控制器每次看“现在离目标差多少”，再决定制冷指令增加还是减少。" if beginner else "所有控制器都放在同一个闭环环境中比较：3R2C 热模型加执行器延迟/惯性，Modelica 负责物理参考，FOPDT 负责控制整定代理。")
    heading(doc,"1.1 3R2C 热模型：两个蓄热体和三条换热通道",2)
    para(doc,"室内空气/设备是响应快的热容，墙体/机柜是响应慢的热容；室外到室内、室外到墙体、墙体到室内是三条热阻通道。")
    equation(doc,[("sub","C","z"),"·",("frac",["d",("sub","T","z")],["dt"])," = ",("frac",[("sub","T","o"),"−",("sub","T","z")],[("sub","R","oz")])," + ",("frac",[("sub","T","w"),"−",("sub","T","z")],[("sub","R","zw")])," + ",("sub","Q","int")," − ",("sub","Q","cool")],"室内空气能量平衡方程")
    equation(doc,[("sub","C","w"),"·",("frac",["d",("sub","T","w")],["dt"])," = ",("frac",[("sub","T","o"),"−",("sub","T","w")],[("sub","R","ow")])," + ",("frac",[("sub","T","z"),"−",("sub","T","w")],[("sub","R","zw")])],"墙体/机柜能量平衡方程")
    equation(doc,[("sub","Q","cool")," = actuator(","u",")·",("sub","Q","max")],"制冷量与执行器状态")
    code(doc,"# 离散实现：5 min dead time + 4 min first-order lag","ThermalPlant3R2C.step() 中的执行器实现")
    table(doc,["参数","名义值","单位","含义"],[
        ["Cz","1.8e6","J/K","室内空气/设备快热容"],["Cw","12.0e6","J/K","墙体/机柜慢热容"],
        ["Roz","0.012","K/W","室外到室内热阻"],["Rzw","0.006","K/W","室内与墙体热阻"],
        ["Row","0.020","K/W","室外到墙体热阻"],["Qmax","7000","W","最大制冷能力"],
        ["dt","1","min","快速代理步长"],["u","0-1","—","压缩机容量/PWM归一化指令"]],[1700,1900,1200,4560],8.8)
    heading(doc,"1.2 统一输入、输出和 PI 内环",2)
    table(doc,["位置","输入","输出","更新"],[
        ["场景生成器","To、Tsp、内部/开门热负荷","时间序列","1 min"],["AI整定器","e=Tz-Tsp、Δe、可选工况","Kp、Ki","2 s"],
        ["PI内环","e、Kp、Ki、积分状态","u=0-1","目标100 ms"],["热对象","u、To、Qint","Tz、Tw、Qcool","1 min代理"],
    ],[1800,3000,2700,1860],8.6)
    equation(doc,[("sub","u","raw")," = ",("sub","K","p"),"·e + ",("sub","K","i"),"·",("int","e dt"),"；  u = clip(",("sub","u","raw"),", 0, 1)"],"并联 PI 控制律")
    para(doc,"积分采用条件抗饱和：输出饱和且误差仍把输出推向边界时停止积分。")
    heading(doc,"1.3 三类场景",2)
    table(doc,["场景","初始/设定","扰动","时长"],[
        ["初次快速降温","31.5°C -> 24°C；To=35°C","室外±2°C；内部650 W","4 h"],
        ["设定温度突变","25°C在1 h变23°C；Tz=25.5°C","室外±3°C；内部800 W","5 h"],
        ["持续外界热扰动","Tz=25°C；Tsp=24°C；To=37°C","室外±4.5°C；人员+1800 W；第3 h开门12 min +2600 W","6 h"],
    ],[2100,3100,3100,1060],8.4)
    heading(doc,"1.4 FOPDT：控制器使用的简化地图",2)
    para(doc,"FOPDT 不是更真实的房间，而是把阶跃响应压缩成三个控制参数：K=指令改变后的最终温降比例，L=开始响应前的等待，tau=响应有多慢。")
    equation(doc,["G(s) = −",("frac",["K·",("sup","e","−Ls")],["τs + 1"])],"FOPDT 传递函数")
    equation(doc,[("sub","T","hat"),"(t) = ",("sub","T","0")," − A[1 − ",("sup","e","−(t−L)/τ"),"]，  t > L"],"FOPDT 阶跃响应")
    equation(doc,["K = ",("frac","A","Δu")],"过程增益")
    table(doc,["量","本次结果","来源/意义"],[
        ["delta_u","0.12","名义制冷指令增加12个百分点"],["A","6.61978°C","168 h 阶跃拟合的最终温降"],
        ["K","55.164866 °C/指令","A/delta_u"],["tau","911.080121 min","有界最小二乘联合拟合"],
        ["L","5 min","执行器延迟与拟合结果"],["RMSE","0.381828°C","44次残差函数评价后的全曲线误差"],
        ["诊断交点","t28=104 min；t63=786 min","仅作诊断，不把截尾末点当 t63"],
    ],[1800,2200,5360],8.5)
    heading(doc,"1.5 由 FOPDT 得到 Z-N 与 IMC 的 Kp、Ki",2)
    equation(doc,[("sup",[("sub","K","p")],"ZN")," = 0.9",("frac","τ","KL")],"Ziegler-Nichols 比例增益")
    equation(doc,["T",("sub","i","ZN")," = 3.33L；  ",("sup",[("sub","K","i")],"ZN")," = ",("frac",[("sup",[("sub","K","p")],"ZN")],["T",("sub","i","ZN")])],"Ziegler-Nichols 积分参数")
    equation(doc,["λ = max(τ/3, 3L, 12 min)；  ",("sup",[("sub","K","p")],"IMC")," = ",("frac","τ","K(λ+L)")],"IMC 闭环速度与比例增益")
    equation(doc,["T",("sub","i","IMC")," = min[τ, 4(λ+L)]；  ",("sup",[("sub","K","i")],"IMC")," = ",("frac",[("sup",[("sub","K","p")],"IMC")],["T",("sub","i","IMC")])],"IMC 积分参数")
    callout(doc,"重要纠正","旧版12 h试验只达到最终温降约59.3%，没有越过63.2%，却把末点当t63。当前版本用168 h长阶跃+有界最小二乘，并保存每个候选参数和RMSE。","FFF4E5")

def bayes(doc, beginner=False):
    heading(doc,"2. 贝叶斯优化自动整定：从试参数到固定 Kp、Ki")
    para(doc,"贝叶斯优化就是“会记账的试错”：每次完整仿真一组旋钮，记下效果，再用高斯过程预测下一次最值得试的位置。最终只冻结一套参数，所以它是自动整定，不是在线自整定。" if beginner else "BO 在运行前搜索全局固定 PI 增益；每个候选点都完成完整场景 rollout 后才计入目标函数。")
    heading(doc,"2.1 输入和目标",2)
    table(doc,["项目","本次设置","原因"],[
        ["变量","Kp、Ki","输出固定PI增益"],["边界","Kp[0.002,1.5]；Ki[1e-5,0.08]","安全可部署范围"],
        ["坐标","log(Kp)、log(Ki)","Ki跨数量级，避免线性采样浪费"],["训练工况","20个随机场景，每个5 h/1 min","覆盖热容、热阻、容量和延迟变化"],
        ["一次评估","完整rollout+objective","不能只看一个时刻"],
    ],[1800,3500,4060],8.7)
    equation(doc,["J = 2·IAE + 1.5·ITAE + 10·",("sub","V","comfort")," + 3·",("sub","ΔT","under")," + 0.35·",("sub","t","settle")," + 0.08·",("sub","M","u")," + 0.7·Var(u) + 0.015·",("sub","E","cool")],"统一评价目标：J 越小越好")
    heading(doc,"2.2 输入到输出的12步",2)
    for s in [
        "生成随机训练工况，输出To、Tsp、Qint和3R2C参数。",
        "做长阶跃辨识得到FOPDT的K、tau、L。",
        "建立Kp、Ki安全边界，并编码到log坐标的[0,1]^2。",
        "放入IMC初始点和随机探索点；单场景标签生成还加入Z-N初始点。",
        "对每个点完整仿真，得到Tz、误差、u和能耗。",
        "把时间序列压缩为objective，越小越好。",
        "用已有(candidate, objective)拟合Matern+WhiteKernel高斯过程。",
        "在512个随机候选点预测均值mu和不确定度sigma。",
        "计算EI：预测可能更好或不确定度更大都值得尝试。",
        "选择EI最大的点，实际rollout后加入样本集。",
        "重复4次EI迭代，取objective最小的候选。",
        "输出完整history并冻结最终Kp、Ki；运行期不再搜索。",
    ]: step(doc,s)
    heading(doc,"2.3 每个候选怎样改变",2)
    b=rows("outputs/bayesian_search_history.csv")
    data=[[r["evaluation"],r["phase"],num(r["candidate_kp"],6),num(r["candidate_ki"],6),num(r["objective"],5),num(r["best_kp"],6),num(r["best_ki"],6),num(r["best_objective"],5),num(r["expected_improvement"],4)] for r in b]
    table(doc,["评估","阶段","候选Kp","候选Ki","目标","当前最优Kp","当前最优Ki","最优目标","EI"],data,[650,1500,1100,1100,1100,1100,1100,1100,610],7.3)
    source(doc,"outputs/bayesian_search_history.csv；初始点的EI为空，因为高斯过程尚未拟合。")
    picture(doc,"outputs/bayesian_search_trace.png","贝叶斯搜索轨迹：第10次评估得到最终最优点。")
    table(doc,["最终输出","数值","含义"],[["固定Kp","0.8710247818","运行期所有场景使用同一套比例增益"],["固定Ki","0.00704244314 min^-1","运行期所有场景使用同一套积分增益"],["最优目标","40.75204078","全局训练工况平均综合目标"],["评价次数","11","7个初始点+4次EI迭代"]],[1800,2400,5160],8.8)
    heading(doc,"2.4 rollout和耗时",2)
    para(doc,"本次quick整条流水线耗时约57.1 s（源码记录总耗时，没有把每个模块单独计时，因此不能把57.1 s冒充三个算法各自耗时）。可审计的计算量是：标签BO 20*(7初始+3 EI)=200次rollout；全局BO 11个候选*20个训练场景=220次rollout；另外每个场景跑Z-N和IMC各一次，共40次。")
    code(doc,"model=identify_fopdt(scenario)\nseeds=[encode(imc_pi(model)), encode(ziegler_nichols_pi(model))]\ngp.fit(x,y)\nmean,std=gp.predict(candidates,return_std=True)\nei=improvement*norm.cdf(z)+std*norm.pdf(z)\nnext_point=candidates[argmax(ei)]\nreturn TuneResult(kp=best_kp,ki=best_ki,score=best_score)","BayesianGainTuner.tune 核心代码")
    callout(doc,"优点与边界","样本效率高、能利用IMC/Z-N先验、轨迹可审计；但每个候选都要完整仿真，且固定增益遇到工况迁移不会主动变化。")

def fnn(doc, beginner=False):
    heading(doc,"3. FNN在线自整定：把离线试验压缩成25条规则")
    para(doc,"这里的FNN不是大型深度网络，而是一张小经验表：输入e和delta_e各分5个区间，25个格子各存一组Kp、Ki；运行时只计算附近4个格子。" if beginner else "实现是轻量零阶TSK规则面：BO离线标签拟合5x5规则后件，在线用双线性插值激活4条规则，并交给安全PI内环。")
    heading(doc,"3.1 结构",2)
    table(doc,["层","实现","参数变化"],[
        ["输入e","中心[-5,-2.5,0,2.5,5]°C","误差大时倾向增大Kp"],["输入delta_e","中心[-1,-0.5,0,0.5,1]°C/步","判断升温、稳定、降温"],
        ["规则","5x5=25条；每条保存Kp/Ki","训练后形成rule_table.npy"],["训练目标","20个场景的BO label","在log空间累计平均"],
        ["在线","searchsorted+双线性插值","每次只算4条，每2 s更新"],["安全","边界+25%变化率+IMC回退","非有限值立即回退"],
    ],[1800,4000,3560],8.5)
    heading(doc,"3.2 从BO标签到动态增益的10步",2)
    for s in [
        "先对20个虚拟工况做BO，得到label_kp、label_ki。",
        "每个场景replay一次，收集每分钟(e,delta_e)，共20*301=6020个状态样本。",
        "将e和delta_e分别映射到5个中心，得到规则索引。",
        "在log空间累加规则目标log_sum[i,j]。",
        "累计counts[i,j]，表示规则被多少状态访问。",
        "每加入一个场景，用log_sum/counts更新已占用格子。",
        "未占用格子保留IMC fallback，避免空白区输出危险参数。",
        "exp回实际空间，得到5x5x2 rule_table。",
        "计算log RMSE、覆盖率、最大规则变化，记录每批快照。",
        "在线用4个相邻规则插值，再经过边界和25%变化率限制输出Kp、Ki。",
    ]: step(doc,s)
    equation(doc,[("sub","K","p")," = Σ",("sub","i,j","N(e,Δe)")," w",("sub","ij",""),("sub","K","p,ij"),"；  ",("sub","K","i")," = Σ w",("sub","ij",""),("sub","K","i,ij")],"FNN 双线性插值：只激活相邻4条规则")
    code(doc,"e0,e1,ew=_active(error_c,centers)\nd0,d1,dw=_active(delta_error_c,delta_centers)\nfor ei,we in ((e0,1-ew),(e1,ew)):\n  for di,wd in ((d0,1-dw),(d1,dw)):\n    rule_kp,rule_ki=rule_table[ei,di]\n    kp += we*wd*rule_kp\n    ki += we*wd*rule_ki\nreturn kp,ki","FNNGainController._propose：25条规则中每次只计算4条")
    heading(doc,"3.3 20个训练场景中规则表如何变化",2)
    fn=rows("outputs/fnn_training_history.csv")
    data=[[r["training_scenario"],r["state_samples"],r["occupied_rules"],r["rule_coverage_pct"],num(r["training_log_rmse"],5),num(r["max_rule_log_change"],5),num(r["mean_rule_kp"],5),num(r["mean_rule_ki"],6),num(r["center_rule_kp"],5),num(r["center_rule_ki"],6),num(r["high_error_rule_kp"],5),num(r["high_error_rule_ki"],6)] for r in fn]
    table(doc,["场景","样本","占用","覆盖%","logRMSE","最大变化","平均Kp","平均Ki","中心Kp","中心Ki","高误差Kp","高误差Ki"],data,[560,700,700,700,800,950,760,800,760,800,770,850],6.8)
    source(doc,"outputs/fnn_training_history.csv；最终覆盖5/25条规则，未覆盖格子继续使用IMC fallback。")
    picture(doc,"outputs/fnn_training_trace.png","FNN规则覆盖率、误差与参数变化。")
    heading(doc,"3.4 运行期Kp、Ki",2)
    on=rows("outputs/online_gain_history.csv")
    data=[[r["minute"],num(r["kp"],7),num(r["ki"],7),"是" if r["fallback_ood"]=="1" else "否"] for r in on]
    table(doc,["时间min","Kp","Ki min^-1","越界回退"],data,[1700,2400,2700,2560],8.1)
    source(doc,"outputs/online_gain_history.csv；文件以30 min采样展示轨迹，SafeAdaptivePI实际AI更新周期是2 s。")
    picture(doc,"outputs/dynamic_comparison.png","动态工况下温度、压缩机指令和在线增益。")
    table(doc,["量","数值","工程含义"],[["状态样本","6020","20场景*301步"],["最终覆盖","5/25=20%","quick训练尚未覆盖全规则"],["最终logRMSE","1.06325","BO标签的对数拟合误差"],["案例AI推理","约60-72 us/次","PC软件在环参考，不是目标板实测"],["在线更新","每2 s","PI内环仍为100 ms"]],[1900,2400,5060],8.5)
    callout(doc,"优点与边界","FNN推理轻、规则可解释、适合MCU；但当前是稀疏BO标签压缩，不是端到端深度网络。覆盖率和logRMSE说明还应增加训练工况。")

def rl(doc, beginner=False):
    heading(doc,"4. 增量式强化学习：从rollout、奖励到Q表")
    para(doc,"RL像在虚拟房间里玩游戏：当前状态下把旋钮小幅移动，下一步温度更好就奖励，差就扣分。真实运行时不再随机探索。" if beginner else "采用安全约束的增量式Q-learning。动作只改变Kp/Ki缩放，不直接命令压缩机；部署时冻结argmax策略，所有输出仍经过安全层。")
    heading(doc,"4.1 状态、动作和超参数",2)
    table(doc,["对象","设置","原因"],[
        ["状态","e、delta_e各5档，共25状态","MCU查表且可解释"],["动作","9个增量：-10%、0、+10%，另含Kp+20%","限制单步调参幅度"],
        ["Q表","5x5x9","每个状态-动作的长期价值"],["alpha","0.12","每次TD更新移动12%"],["gamma","0.94","重视后续温度效果"],
        ["epsilon","max(0.03,0.25*(1-episode/500))","从0.25降到0.03"],["训练","500回合*60步=30000环境步","纯虚拟rollout"],
    ],[1800,3600,3960],8.6)
    heading(doc,"4.2 一次rollout的11步",2)
    for s in [
        "随机抽场景，随机初始化Tz在Tsp-5到Tsp+9°C；Tw=0.75*Tz+0.25*To。",
        "Q表全0，kp_scale=ki_scale=1；第一回合无先验偏好。",
        "得到误差e，并随机生成第一步delta_e，覆盖粗网格状态。",
        "把e、delta_e分箱到25个状态，更新visited_states。",
        "epsilon-greedy选9个动作：随机探索或当前Q最大动作。",
        "更新kp_scale、ki_scale，并裁剪在0.5到1.8。",
        "Kp=clip(Kp_IMC*kp_scale,0.002,1.5)，Ki=clip(Ki_IMC*ki_scale,1e-5,0.08)。",
        "用PI和当前参数推进延迟/惯性3R2C，得到next_error。",
        "奖励r=-abs(next_error)-0.18*动作幅度-0.04*abs(kp_scale-1)-0.03*u。",
        "target=r+gamma*max(Q(next_state))；Q=Q+alpha*(target-Q)。",
        "回合结束记录reward、20回合均值、TD、覆盖率、策略变化和Kp/Ki；500回合后冻结argmax。",
    ]: step(doc,s)
    equation(doc,["r = −|e'| − 0.18·",("sub","||a||","1")," − 0.04|",("sub","s","p"),"−1| − 0.03u"],"RL 即时奖励")
    equation(doc,["target = r + γ ",("sub","max","a")," Q(s',a)；  Q(s,a) ← Q(s,a) + α[target − Q(s,a)]"],"RL TD 目标与 Q-learning 更新")
    code(doc,"action=epsilon_random_or_argmax(q[state])\nkp_scale=clip(kp_scale*(1+gain_move[0]),0.5,1.8)\nki_scale=clip(ki_scale*(1+gain_move[1]),0.5,1.8)\ncontroller.kp=clip(fallback_kp*kp_scale,0.002,1.5)\ncontroller.ki=clip(fallback_ki*ki_scale,1e-5,0.08)\nreward=-(abs(next_error)+0.18*abs(gain_move).sum()+0.04*abs(kp_scale-1)+0.03*command)\ntarget=reward+gamma*max(q[next_state])\nq[state+(action,)]+=alpha*(target-q[state+(action,)])","train_offline_q_policy 核心更新")
    heading(doc,"4.3 训练轨迹和参数变化",2)
    rr=rows("outputs/rl_training_history.csv"); data=[]
    for r in rr:
        ep=int(float(r["episode"]))
        if ep<=10 or ep%25==0 or ep>=491:
            data.append([r["episode"],r["environment_steps"],num(r["epsilon"],4),num(r["episode_reward"],6),num(r["moving_average_reward_20"],6),num(r["mean_abs_td_error"],5),r["visited_states"],r["state_coverage_pct"],r["greedy_policy_changes"],num(r["mean_kp"],6),num(r["mean_ki"],7),num(r["final_kp"],6),num(r["final_ki"],7)])
    table(doc,["回合","累计步","eps","奖励","20回合均值","平均|TD|","状态","覆盖%","策略变化","平均Kp","平均Ki","末Kp","末Ki"],data,[500,700,620,850,900,720,700,700,760,760,800,700,850],6.4)
    source(doc,"outputs/rl_training_history.csv；表展示首10、每25回合和末10回合，CSV保存全部500回合。")
    picture(doc,"outputs/rl_training_trace.png","RL奖励、TD误差、状态覆盖率和增益轨迹。")
    table(doc,["最终观测","数值","解释"],[["回合","500","每回合最多60步，共30000环境步"],["eps","0.03","训练期最低探索；部署时关闭随机探索"],["覆盖","25/25=100%","粗网格状态全部访问"],["末20回合平均奖励","-227.7505","仍有明显波动"],["末回合平均|TD|","1.5611","Q值仍在调整"],["末回合Kp/Ki","0.0963029 / 4.9966e-05","最后一个rollout的参数"]],[2200,2400,4760],8.6)
    heading(doc,"4.4 训练和部署的差别",2)
    table(doc,["阶段","随机吗","更新Q吗","参数变化","安全措施"],[
        ["训练","epsilon-greedy","是","虚拟环境按动作改变","动作、缩放、边界均裁剪"],
        ["部署","否，argmax","否","每2 s提出目标，经25%变化率限制","NaN/越界/异常回退IMC"],
    ],[1300,2000,1300,2200,2560],8.2)
    callout(doc,"优点与边界","RL可以学习状态相关策略，但训练成本高。当前已访问25/25状态，奖励仍未稳定，所以应作为离线原型继续训练，而不是宣称已经收敛。","FFF4E5")

def comparison(doc, beginner=False):
    heading(doc,"5. 如何比较优劣：不只看一个最终数字")
    para(doc,"没有一个算法在所有场景都最好：贝叶斯适合提前调好后稳定运行；FNN是轻量动态经验表；RL表达能力强但当前训练不足；IMC平稳却保守；Z-N快但激进，必须安全限幅。")
    sm=rows("outputs/holdout_summary.csv")
    data=[[r["controller"],r["scenarios"],num(r["mean_objective"],5),num(r["mean_rmse_c"],4),num(r["mean_itae_c_hour2"],4),num(r["mean_settling_time_hour"],4),num(r["mean_cooling_energy_kwh"],4),num(r["mean_control_movement"],4),num(r["mean_compressor_output_variance"],5),num(r["mean_mean_ai_inference_us"],5)] for r in sm]
    table(doc,["控制器","场景数","平均目标","RMSE","ITAE","调节h","能耗kWh","动作量","指令方差","AI us"],data,[1750,750,900,850,850,850,950,900,950,560],7.0)
    source(doc,"outputs/holdout_summary.csv；6个随机留出工况，与三类案例互补。")
    table(doc,["场景","优先观察","当前工程解释"],[
        ["初次快速降温","调节时间、过冷、压缩机动作","Z-N/贝叶斯/FNN较快；IMC过保守；RL波动更大。"],
        ["设定温度突变","ITAE、过冷、恢复","贝叶斯/FNN误差积分低；RL受未完全收敛影响。"],
        ["持续热扰动","开门恢复、参数轨迹","FNN能随e、delta_e变；固定BO不会主动改；IMC恢复慢。"],
    ],[1900,3100,4360],8.5)
    picture(doc,"outputs/training_convergence_overview.png","三种AI整定/训练过程的收敛概览。")
    heading(doc,"5.1 MCU落地检查点",2)
    table(doc,["项目","当前证据","量产前补充"],[
        ["调度","PI 100 ms；AI 2 s；AI calls=20","目标板定时器和最坏周期实测"],["RAM/ROM","PC ABI: SafePI 32 B；对象68 B","目标编译器map文件确认"],
        ["推理","FNN 4规则；RL 25x9 Q表","定点误差、Flash对齐、Wokwi/目标板留证"],
        ["安全","u边界、Kp/Ki边界、变化率、IMC回退","传感器断线、NaN、通信丢包、压缩机故障"],
    ],[1700,3500,4160],8.3)
    callout(doc,"工程结论","当前系统已把“自动整定”和“在线自整定”严格分开：BO输出固定参数，FNN/RL输出动态参数；quick版本证明流程可运行，但FNN覆盖率和RL收敛性仍不足以宣称AI必然优于基线。","EEF6FF")

def appendix(doc):
    heading(doc,"附录：机器可读审计入口")
    table(doc,["文件","记录内容","对应链路"],[
        ["outputs/bayesian_search_history.csv","11次候选、目标、EI、当前最优","BO搜索"],
        ["outputs/fnn_training_history.csv","20批规则表快照","FNN拟合"],
        ["outputs/online_gain_history.csv","运行期Kp/Ki轨迹","在线自整定"],
        ["outputs/rl_training_history.csv","500回合reward、eps、TD、覆盖、增益","RL训练"],
        ["outputs/fnn_rule_table.npy","最终5x5x2规则后件","FNN部署表"],
        ["outputs/rl_q_table.npy","最终5x5x9 Q表","RL冻结策略"],
        ["outputs/case_metrics.csv","三类场景五种控制器指标","工程评价"],
        ["hvac_pid/tuning.py","BO、EI、FOPDT标签","Bayesian"],
        ["hvac_pid/ai_controllers.py","FNN/RL训练和推理","FNN/RL"],
    ],[3150,3700,2510],8.3)
    heading(doc,"复现命令",2)
    code(doc,"python main.py --quick\npython -m pytest tests -q\nstreamlit run streamlit_app.py","quick复现入口")
    callout(doc,"边界声明","这是软件仿真与PC软件在环验证，不代表真实压缩机或目标MCU最终性能。OpenModelica运行时需本机安装；量产前必须完成目标板编译、测量和保护验证。","FFF4E5")

def build(formal):
    doc=Document(); beginner=not formal
    setup(doc,"HVAC AI-PID | "+("全链路提交报告" if formal else "零基础解释版"),compact=beginner)
    title(doc, "变频精密/机柜空调 AI-PID\\n贝叶斯、FNN 与增量式 RL 全链路技术报告" if formal else "贝叶斯、FNN 与强化学习\\n到底怎样一步步调空调？",
          "从热模型、输入数据到训练/搜索、参数变化、控制输出与工程落地" if formal else "先讲直觉，再对照真实代码和训练轨迹",
          [("用途","课程/项目提交；算法实现、训练过程与结果可追溯" if formal else "面向不了解PID、FOPDT、神经网络和强化学习的读者"),
           ("实验版本","quick：20训练工况、6留出工况、每标签3次BO迭代" if formal else "对应正式报告：01_贝叶斯_FNN_RL_全链路技术报告.docx"),
           ("数据来源","outputs/下HTML、CSV、NPY、PNG、Excel和PPT"),
           ("日期",str(date.today()))],beginner=beginner)
    callout(doc,"一页结论","贝叶斯是提前试参数并冻结一套固定Kp/Ki；FNN用BO标签拟合25条规则，运行时每2 s动态输出；RL用500回合、30000个虚拟环境步学习25x9 Q表，部署时关闭随机探索。三者不能混称为同一种AI自整定。")
    common(doc, beginner)
    doc.add_page_break(); bayes(doc, beginner)
    doc.add_page_break(); fnn(doc, beginner)
    doc.add_page_break(); rl(doc, beginner)
    doc.add_page_break(); comparison(doc, beginner)
    appendix(doc)
    path=OUT/("01_贝叶斯_FNN_RL_全链路技术报告.docx" if formal else "02_贝叶斯_FNN_RL_零基础解释版.docx")
    doc.save(path); return path

if __name__=="__main__":
    print(build(True)); print(build(False))
