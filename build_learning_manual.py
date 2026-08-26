from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs_adaptive_final_v2"
DOC_DIR = OUT / "learning_document"
FIG_DIR = DOC_DIR / "figures"
DOCX_PATH = DOC_DIR / "HVAC_AI_PID_零基础学习手册.docx"
FONT_PATH = Path(r"C:\Windows\Fonts\msyh.ttc")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
PALE_GOLD = "FFF4D6"
PALE_RED = "FDECEC"
GREEN = "E8F5E9"
MUTED = "666666"


def rows(name: str) -> list[dict[str, str]]:
    with (OUT / name).open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def f(value: object, digits: int = 4) -> str:
    try:
        return f"{float(value):.{digits}g}"
    except (TypeError, ValueError):
        return str(value)


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_run(run, size=None, bold=None, color=None, italic=None, math=False) -> None:
    run.font.name = "Cambria Math" if math else "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Cambria Math" if math else "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Cambria Math" if math else "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, fld_end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "HVAC AI-PID 零基础学习手册  |  仿真、训练、嵌入式与验收"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(hp.runs[0], size=8.5, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("第 ")
    set_run(r, size=8.5, color=MUTED)
    add_page_number(fp)
    r = fp.add_run(" 页")
    set_run(r, size=8.5, color=MUTED)


def add_para(doc: Document, text: str, *, bold_lead: str | None = None, align=None) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run(lead, bold=True)
        body = p.add_run(text[len(bold_lead):])
        set_run(body)
    else:
        run = p.add_run(text)
        set_run(run)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run(p.add_run(item))


def add_numbers(doc: Document, items: list[str]) -> None:
    # Write the ordinal explicitly so every independent procedure restarts at 1.
    # Word otherwise keeps one hidden numbering sequence across the whole file.
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        set_run(p.add_run(f"{index}.  "), bold=True)
        set_run(p.add_run(item))


def add_formula(doc: Document, formula: str, explanation: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_fill(cell, PALE_GRAY)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(formula), size=11.5, math=True)
    if explanation:
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run(p2.add_run(explanation), size=9.5, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc: Document, title: str, text: str, fill: str = PALE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    p = cell.paragraphs[0]
    set_run(p.add_run(title + "："), bold=True, color=INK)
    set_run(p.add_run(text), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc: Document, headers: list[str], data: list[list[object]], widths: list[int] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    # Repeat the header on every page and keep it with the first data row.
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_pr.append(repeat_header)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_fill(cell, PALE_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.keep_with_next = True
        set_run(p.add_run(str(header)), bold=True, size=9.5, color=INK)
    for row in data:
        cells = table.add_row().cells
        # A data row may move to the next page, but it must not be split in half.
        row_pr = table.rows[-1]._tr.get_or_add_trPr()
        row_pr.append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            set_run(p.add_run(str(value)), size=9.2)
    if widths is None:
        base = 9360 // len(headers)
        widths = [base] * len(headers)
        widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def font(size: int, bold=False):
    if FONT_PATH.exists():
        return ImageFont.truetype(str(FONT_PATH), size=size)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, box_width: int, ft) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if draw.textbbox((0, 0), candidate, font=ft)[2] <= box_width or not current:
            current = candidate
        else:
            lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def draw_flow(path: Path, title: str, nodes: list[tuple[str, str]], footer: str = "") -> None:
    width = 1500
    box_h = 118
    gap = 48
    height = 130 + len(nodes) * (box_h + gap) + 80
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 40), title, font=font(34, True), fill="#1F4D78")
    x0, x1 = 170, width - 170
    y = 120
    colors = ["#E8EEF5", "#EDF6FF", "#E8F5E9", "#FFF4D6"]
    for index, (name, detail) in enumerate(nodes):
        fill = colors[index % len(colors)]
        draw.rounded_rectangle((x0, y, x1, y + box_h), radius=18, fill=fill, outline="#4A6D8C", width=3)
        draw.text((x0 + 30, y + 18), name, font=font(27, True), fill="#0B2545")
        detail_lines = wrap(draw, detail, x1 - x0 - 330, font(22))
        for li, line in enumerate(detail_lines[:2]):
            draw.text((x0 + 300, y + 18 + li * 34), line, font=font(22), fill="#333333")
        if index < len(nodes) - 1:
            cx = width // 2
            draw.line((cx, y + box_h, cx, y + box_h + gap - 10), fill="#4A6D8C", width=5)
            draw.polygon([(cx - 10, y + box_h + gap - 22), (cx + 10, y + box_h + gap - 22), (cx, y + box_h + gap - 4)], fill="#4A6D8C")
        y += box_h + gap
    if footer:
        draw.text((70, height - 55), footer, font=font(20), fill="#666666")
    image.save(path)


def draw_fopdt(path: Path) -> None:
    width, height = 1500, 800
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), "一阶加延迟阶跃响应：63.2%从哪里来", font=font(34, True), fill="#1F4D78")
    left, top, right, bottom = 150, 150, 1400, 680
    draw.line((left, bottom, right, bottom), fill="#333333", width=3)
    draw.line((left, bottom, left, top), fill="#333333", width=3)
    import math
    L, tau = 0.14, 0.23
    points = []
    for i in range(800):
        x = i / 799
        y = 0.0 if x < L else 1.0 - math.exp(-(x - L) / tau)
        px = left + x * (right - left)
        py = bottom - y * (bottom - top - 35)
        points.append((px, py))
    draw.line(points, fill="#2E74B5", width=6)
    xL = left + L * (right - left)
    x63 = left + (L + tau) * (right - left)
    y63 = bottom - 0.63212056 * (bottom - top - 35)
    draw.line((xL, bottom, xL, top + 70), fill="#999999", width=3)
    draw.line((x63, bottom, x63, y63), fill="#C77700", width=4)
    draw.line((left, y63, x63, y63), fill="#C77700", width=4)
    draw.ellipse((x63 - 9, y63 - 9, x63 + 9, y63 + 9), fill="#C77700")
    draw.text((xL - 15, bottom + 12), "L", font=font(24, True), fill="#555555")
    draw.text((x63 - 55, bottom + 12), "L + τ", font=font(24, True), fill="#C77700")
    draw.text((left - 110, y63 - 18), "63.2%", font=font(24, True), fill="#C77700")
    draw.text((360, 710), "y(t)=A[1-exp(-(t-L)/τ)]；当 t=L+τ 时，1-exp(-1)=0.6321", font=font(25), fill="#0B2545")
    image.save(path)


def draw_coverage(path: Path, old_fnn: int, old_rl: int, final_fnn: int, final_rl: int) -> None:
    width, height = 1500, 780
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), "状态/规则覆盖：整改前与整改后", font=font(34, True), fill="#1F4D78")
    panels = [("旧FNN", old_fnn), ("最终FNN", final_fnn), ("旧RL热状态", old_rl), ("最终RL热状态", final_rl)]
    for panel, (label, count) in enumerate(panels):
        x0 = 70 + panel * 355
        y0 = 155
        draw.text((x0, 100), f"{label}  {count}/25", font=font(25, True), fill="#0B2545")
        for e in range(5):
            for d in range(5):
                idx = e * 5 + d
                fill = "#69B578" if idx < count else "#E2E5E9"
                x = x0 + d * 55
                y = y0 + (4 - e) * 55
                draw.rectangle((x, y, x + 46, y + 46), fill=fill, outline="#FFFFFF", width=2)
        draw.text((x0, 455), "横轴：误差变化率档", font=font(18), fill="#555555")
        draw.text((x0, 485), "纵轴：温度误差档", font=font(18), fill="#555555")
    draw.rounded_rectangle((70, 570, 1430, 710), radius=16, fill="#FFF4D6", outline="#C89A32", width=2)
    note = "覆盖提升来自分层物理工况、统一误差变化率和合理分箱，不是手工把25个格子标为已访问。RL另外自然访问47/75个‘热状态×容量模式’组合；冷房间高制冷等不安全组合不应为了100%而伪造。"
    lines = wrap(draw, note, 1280, font(22))
    for i, line in enumerate(lines):
        draw.text((105, 595 + i * 34), line, font=font(22), fill="#5D4500")
    image.save(path)


def add_figure(doc: Document, path: Path, caption: str, width=6.3) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    # Use the visible caption as alternative text for screen-reader users.
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    set_run(cp.add_run(caption), size=9, color=MUTED, italic=True)


def chapter(doc: Document, title: str, page_break_before: bool = False) -> None:
    heading = doc.add_heading(title, level=1)
    heading.paragraph_format.page_break_before = page_break_before


def build() -> Path:
    DOC_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)

    fnn_old = rows("../outputs_review_final/fnn_training_history.csv") if False else []
    fnn_history = rows("fnn_training_history.csv")
    rl_history = rows("rl_training_history.csv")
    summary = rows("holdout_summary.csv")
    imc_rows = rows("imc_lambda_tuning.csv")
    bo = rows("global_bayesian_tuning.csv")[0]
    fopdt = rows("classical_tuning_history.csv")[0]
    fnn_final = fnn_history[-1]
    rl_final = rl_history[-1]
    selected_imc = next(item for item in imc_rows if int(float(item["selected"])) == 1)
    summary_by_name = {item["controller"]: item for item in summary}

    flow_full = FIG_DIR / "full_pipeline.png"
    draw_flow(flow_full, "从需求到真实空调：完整工程链路", [
        ("1 需求与安全边界", "温度范围、响应时间、最低频率、启停与联锁"),
        ("2 数据与物理模型", "BMS/实验数据；3R2C、执行器延迟、噪声"),
        ("3 基线控制", "FOPDT辨识；Z-N、IMC与固定PI"),
        ("4 离线优化/训练", "BO标签、FNN规则、RL交互经验"),
        ("5 三段切分与验收", "拟合、内部验证早停、独立留出测试"),
        ("6 自动导出", "参数、规则/策略、边界、版本与CRC"),
        ("7 MCU固件", "100 ms PI、秒级AI、执行器限制、安全回退"),
        ("8 SIL/HIL/实机", "数学一致性、开发板、故障注入、受控试运行"),
    ], "本项目目前完成到PC SIL；真实目标板资源实测、HIL和实机尚未完成。")
    control_flow = FIG_DIR / "runtime_loop.png"
    draw_flow(control_flow, "嵌入式运行时闭环", [
        ("传感器输入", "室温、设定值、压力/排温、变频器状态"),
        ("过滤与诊断", "去噪、断线/越界判断、时间戳"),
        ("FNN/RL外环", "读取e、ė和容量模式，提出Kp/Ki目标"),
        ("安全增益层", "验收标志、覆盖掩码、±10%、上下限、IMC回退"),
        ("PI内环", "u_req=Kp·e+Ki·∫e dt；100 ms执行"),
        ("压缩机限制器", "0/25%、1%量化、5%/min、5/3 min启停"),
        ("执行器输出", "Modbus/0-10V/PWM目标频率与状态遥测"),
        ("房间响应", "压缩机制冷改变温度，再被传感器测量"),
    ])
    fopdt_fig = FIG_DIR / "fopdt_632.png"
    draw_fopdt(fopdt_fig)
    coverage_fig = FIG_DIR / "coverage_before_after.png"
    draw_coverage(coverage_fig, 12, 16, int(float(fnn_final["occupied_rules"])), int(float(rl_final["visited_states"])))

    doc = Document()
    configure_document(doc)

    # Editorial cover pattern with compact-reference-guide body tokens.
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("工程学习手册"), size=12, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run("HVAC AI-PID"), size=30, bold=True, color=INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("从3R2C、FOPDT和PI，到BO、FNN、RL与嵌入式部署"), size=15, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    set_run(p.add_run("面向完全不了解嵌入式、PID和机器学习的读者"), size=11, italic=True, color=MUTED)
    for _ in range(7):
        doc.add_paragraph()
    add_callout(doc, "本手册的诚实结论", "整改后FNN与RL均覆盖25/25个热状态；但FNN候选在内部验证上没有优于IMC，因此实际部署自动回退。RL通过内部验收，在16个独立工况的平均综合目标上仅比IMC改善约0.37%，属于‘流程有效、收益很小’，不能宣称AI显著优越。", PALE_GOLD)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(f"实验版本：outputs_adaptive_final_v2  |  生成日期：{date.today().isoformat()}"), size=9.5, color=MUTED)
    doc.add_page_break()

    chapter(doc, "阅读地图：先看懂六个最重要的词")
    add_table(doc, ["词", "零基础解释", "本项目中的具体对象"], [
        ["被控对象", "你想改变的真实东西", "机房/机柜空气、墙体和热负荷"],
        ["控制器", "根据误差决定空调该出多少力的程序", "PI，以及调整PI参数的FNN/RL外环"],
        ["模型", "真实系统的数学替身", "3R2C热模型、FOPDT低阶代理"],
        ["训练", "用已知试验反复修改规则或价值表", "FNN拟合25条规则；RL更新Q表"],
        ["嵌入式", "把算法放进资源有限、必须准时工作的MCU", "STM32/ESP32上的C++固件"],
        ["验收", "用没参与训练的数据证明它没有变差", "内部验证门+16个独立留出工况"],
    ], [1300, 3900, 4160])
    add_callout(doc, "建议阅读顺序", "第一次阅读先看第1、4、11、14、16章；第二遍再看公式推导；准备做真实硬件时重点看第17、18章。")

    chapter(doc, "1. 整个项目究竟在解决什么问题")
    add_para(doc, "目标不是让AI直接控制压缩机，而是建立一个分层控制系统：最底层PI快速计算容量请求；FNN或RL较慢地建议Kp、Ki；无论AI建议什么，都必须经过安全限幅、执行器约束和IMC回退。")
    add_figure(doc, flow_full, "图1  完整工程链路和本项目当前位置")
    doc.add_heading("1.1 输入、输出和不可越过的边界", level=2)
    add_table(doc, ["层", "输入", "处理", "输出"], [
        ["物理系统", "室外温度、热负荷、制冷量", "热量流动与蓄积", "室温、墙温"],
        ["传感器层", "真实温度/压力/状态", "采样、滤波、诊断", "可信测量值+故障标志"],
        ["AI外环", "e、ė、实际容量模式", "FNN插值或RL查表", "候选Kp、Ki"],
        ["PI内环", "e、Kp、Ki、dt", "比例+积分+抗饱和", "0～1容量请求"],
        ["执行器安全层", "容量请求、启停状态和计时", "最低频率、量化、斜率、驻留", "可执行容量命令"],
        ["硬件输出", "可执行命令", "通信/PWM/模拟量", "变频器频率或启停许可"],
    ], [1200, 2600, 3150, 2410])

    chapter(doc, "2. 嵌入式开发全流程：算法只是中间一环")
    add_para(doc, "嵌入式控制器不是‘把Python复制到开发板’。它同时包含硬件、驱动、实时调度、故障保护、参数管理、通信和量产验证。")
    add_numbers(doc, [
        "定义需求：允许温度误差、最大响应时间、传感器范围、压缩机允许频率、故障时动作。输出是可测试的需求规格。",
        "选择硬件：MCU、ADC精度、PWM/0-10V/RS-485/CAN、隔离电源、存储和看门狗。输出是原理图、BOM和引脚表。",
        "建立模型与数据协议：决定需要记录的时间戳、温度、设定值、实际频率、压力、报警和负荷。输出是数据字典。",
        "开发控制算法：先固定PI和安全限制，再做BO/FNN/RL。输出是算法代码与离线产物。",
        "固件移植：把浮点、表格和调度变成确定性的C/C++，禁止动态内存和不可控阻塞。输出是可烧录固件。",
        "SIL软件在环：同样输入喂给Python和C++，逐点比较Kp、Ki、请求和限制后命令。",
        "开发板联调：验证ADC、PWM、通信、Flash、RAM、栈和最坏执行时间WCET。",
        "HIL硬件在环：真实开发板接实时仿真器，注入传感器断线、通信超时、压力报警和电源复位。",
        "实机受控试运行：先影子模式，再限制权限，最后逐步扩大范围；所有硬联锁独立于AI。",
        "量产与维护：EMC、环境、老化、标定、升级、CRC、版本回滚和现场日志。",
    ])
    add_callout(doc, "当前项目位置", "Python算法、仿真数据、自动导出和PC SIL已完成；Wokwi仅有编译/启动级证据；真实开发板资源和WCET、HIL、真实变频器与压缩机试验尚未完成。", PALE_GOLD)

    chapter(doc, "3. 房间为什么可以写成热模型")
    doc.add_heading("3.1 能量守恒是所有热模型的起点", level=2)
    add_formula(doc, "C · dT/dt = Q进入 - Q离开", "C是等效热容量，单位J/K；右侧单位W=J/s，所以除以C后得到K/s。")
    add_para(doc, "热容量C越大，同样的净热量造成的温度变化越慢。热阻R描述热量穿过墙体或空气通道的困难程度，两个节点温差为ΔT时，热流近似为ΔT/R。")
    doc.add_heading("3.2 3R2C每个参数的意义", level=2)
    add_table(doc, ["参数", "单位", "物理意义", "数值变大时"], [
        ["Cz", "J/K", "空气及快速热质的等效热容量", "室温短期变化更慢"],
        ["Cw", "J/K", "墙体/机柜等慢热质热容量", "长尾更慢、储热更多"],
        ["Roz", "K/W", "室外到室内空气的直接热阻", "外界对室温影响减弱"],
        ["Rzw", "K/W", "室内空气与慢热质之间热阻", "空气和墙体交换更慢"],
        ["Row", "K/W", "室外到慢热质热阻", "墙体受室外影响减弱"],
        ["Qmax", "W", "满容量制冷量", "同一指令制冷更强"],
    ], [1100, 1100, 3900, 3260])
    add_formula(doc, "Cz·dTz/dt = (To-Tz)/Roz + (Tw-Tz)/Rzw + Qload - Qcool", "室内空气节点的能量平衡。")
    add_formula(doc, "Cw·dTw/dt = (Tz-Tw)/Rzw + (To-Tw)/Row", "墙体/机柜慢节点的能量平衡。")
    add_callout(doc, "模型边界", "3R2C比单一温度模型真实，但仍不是制冷剂回路、换热器结霜、风量分布和压缩机效率图的高保真模型。最终必须用BMS/实验数据标定，并在HIL/实机中验证。")

    chapter(doc, "4. 执行器约束：为什么算法不能直接输出0～100%")
    add_para(doc, "真实变频压缩机通常不是任意0～100%连续设备。0代表停机；一旦运行，可能至少需要20%～30%频率。频率指令还有量化、斜率和最小开停机时间。忽略它们会让仿真中的‘优秀控制器’在真实设备上频繁启停或根本无法执行。")
    add_table(doc, ["约束", "最终项目值", "原因", "实现位置"], [
        ["最低运行容量", "25%", "低于稳定频率可能失油、效率差或不能稳定运行", "CompressorCommandLimiter"],
        ["量化", "1%", "变频器寄存器/命令有有限台阶", "输出前量化"],
        ["运行斜率", "5%/min", "保护机械、电流和制冷回路", "连续内部斜坡"],
        ["最小开机", "5 min", "防止刚启动就停机", "状态计时器"],
        ["最小停机", "3 min", "均压和防止短循环", "状态计时器"],
        ["增益变化", "单次±10%", "防止AI让闭环特性瞬间改变", "SafeAdaptivePI/SafePI"],
    ], [1500, 1400, 3600, 2860])
    add_callout(doc, "100 ms量化陷阱", "如果每100 ms先把5%/min换算成0.000083，再立即四舍五入到1%，命令会永远停在原地。最终C++保留未量化的连续内部斜坡，只把对外输出量化，因此小增量会累积到下一台阶。", PALE_GOLD)

    chapter(doc, "5. 从复杂3R2C到FOPDT：为什么要做阶跃辨识")
    add_para(doc, "Z-N和IMC不需要知道每一堵墙的热容量，只需要一个能概括主要动态的低阶模型。FOPDT代表‘一阶惯性+纯延迟’，用三个量描述：过程增益K、时间常数τ和延迟L。")
    add_formula(doc, "G(s) = -K·exp(-Ls)/(τs+1)", "负号表示制冷指令增大时室温下降。")
    doc.add_heading("5.1 一阶阶跃公式如何从微分方程得到", level=2)
    add_formula(doc, "τ·dy/dt + y = K·Δu")
    add_para(doc, "阶跃后Δu为常数，最终稳态满足dy/dt=0，所以y∞=KΔu。令z=y-y∞，得到τ·dz/dt+z=0；变量分离后dz/z=-dt/τ；积分得到z=C·exp(-t/τ)。用初值y(0)=0求出C=-y∞。")
    add_formula(doc, "y(t)=KΔu·[1-exp(-t/τ)]")
    add_para(doc, "加入纯延迟L，就是在t<L时输出不变，在t≥L时把时间替换成t-L。定义最终温降幅度A=KΔu，就得到项目使用的拟合曲线。")
    add_formula(doc, "T̂(t)=T0-A·[1-exp(-(t-L)/τ)],  t≥L")
    add_figure(doc, fopdt_fig, "图2  一阶响应、延迟L、时间常数τ和63.2%")
    doc.add_heading("5.2 为什么是63.2%", level=2)
    add_formula(doc, "1-exp(-1)=0.63212056")
    add_para(doc, "从延迟结束起经过一个τ，响应完成63.2%；经过2τ完成86.5%，3τ完成95.0%，5τ完成99.3%。它是指数解的数学性质，不是人为经验阈值。")
    doc.add_heading("5.3 为什么最终用整条曲线拟合", level=2)
    add_para(doc, "旧版12 h试验没有到达63.2%，却把末点当成t63，属于截尾偏差。当前代码模拟168 h，并用有界最小二乘同时寻找A、τ、L，使所有采样点的平方误差总和最小。t28/t63只检查曲线是否走过关键区域。")
    add_formula(doc, "min(A,τ,L) Σk [T̂(tk;A,τ,L)-Tmeas(tk)]²")
    add_table(doc, ["本次辨识量", "数值", "解释"], [
        ["K", f(fopdt["process_gain_c_per_u"], 6) + " °C/指令", "容量指令增加1.0对应的拟合稳态温降"],
        ["τ", f(fopdt["time_constant_minutes"], 6) + " min", "主要惯性快慢"],
        ["L", f(fopdt["delay_minutes"], 6) + " min", "纯延迟代理"],
        ["拟合RMSE", f(fopdt["fit_rmse_c"], 5) + " °C", "FOPDT与3R2C阶跃曲线差异"],
        ["辨识时长", f(fopdt["identification_duration_hours"]) + " h", "电脑虚拟时间，不等于实机必须等待7天"],
    ], [1800, 2200, 5360])

    chapter(doc, "6. PI控制器：Kp和Ki到底是什么")
    add_formula(doc, "e(t)=Tzone(t)-Tsetpoint(t)", "制冷系统中e>0表示偏热，需要更多制冷。")
    add_formula(doc, "u(t)=Kp·e(t)+Ki·∫e(t)dt")
    add_para(doc, "比例P只看当前误差：Kp越大，偏热1℃时立即增加的容量越多。积分I累积过去误差：持续偏热会不断增加容量，消除稳态偏差。代码中的dt以分钟计，因此Ki的单位近似为‘每℃·分钟对应的容量’。")
    add_table(doc, ["参数", "太小", "合适", "太大"], [
        ["Kp", "响应慢、扰动后恢复慢", "快速但不过度", "抖动、过冷、频繁触顶"],
        ["Ki", "长期偏差消除慢", "消除稳态偏差", "积分累积、过冲和长时间振荡"],
    ], [1300, 2700, 2660, 2700])
    doc.add_heading("6.1 离散实现", level=2)
    add_formula(doc, "I[k]=I[k-1]+e[k]·Δt；u_raw[k]=Kp·e[k]+Ki·I[k]")
    add_para(doc, "MCU无法连续积分，只能每个采样周期累加。Δt必须使用真实时间单位；若Python用分钟而C++误用秒，Ki会相差60倍。")
    doc.add_heading("6.2 饱和和抗积分饱和", level=2)
    add_para(doc, "当u_raw>1时，压缩机已经满载。若积分仍继续增加，温度接近设定值后仍会长时间满载，导致严重过冷。条件积分只在输出未饱和，或当前误差有助于把输出拉回范围时更新积分。")
    add_formula(doc, "u=clip(u_raw,0,1)")
    add_callout(doc, "PI不是设备安全层", "PI的0～1只是‘请求’。最低频率、斜率、启停、压力和排温联锁必须位于PI之后，并且不能依赖AI。")

    chapter(doc, "7. 如何评价一个控制器")
    add_table(doc, ["指标", "公式/定义", "越小意味着", "局限"], [
        ["IAE", "∫|e|dt", "总体误差更少", "不特别惩罚晚期误差"],
        ["ITAE", "∫t|e|dt", "更早消除误差", "时间越长权重越大"],
        ["RMSE", "sqrt(mean(e²))", "大误差更少", "不说明误差发生时间"],
        ["最大过冷", "max(-e,0)", "设备/舒适风险更低", "只看最坏一点"],
        ["调节时间", "进入±0.5℃并保持60 min", "达到目标更快", "受场景时长影响"],
        ["控制变化量", "Σ|Δu|", "指令更平滑", "不是实际寿命"],
        ["能量代理", "∫Qcool dt", "制冷调用较少", "未含COP、风机和泵"],
    ], [1200, 2350, 2700, 3110])
    add_formula(doc, "J=2·IAE+1.5·ITAE+10·舒适超限+3·最大过冷+0.35·调节时间+0.08·Σ|Δu|+0.7·Var(u)+0.015·能量代理")
    add_para(doc, "J是人为定义的工程折中，不是自然定律。权重改变会改变‘最优’参数，因此必须把公式、单位、权重和场景一起保存。不能只说某算法排名第一而不说明J。")

    chapter(doc, "8. Z-N和IMC：为什么先建立经典基线")
    doc.add_heading("8.1 Z-N反应曲线PI", level=2)
    add_formula(doc, "Kp=0.9τ/(KL)；Ti=3.33L；Ki=Kp/Ti")
    add_para(doc, "Z-N来自经验整定规则，通常响应快但激进。K越大表示对象对指令更敏感，因此Kp应减小；L越大表示延迟越危险，因此Kp也应减小；τ越大表示对象慢，公式允许更大的Kp来推动响应。项目还把结果限制在Kp∈[0.002,1.5]、Ki∈[10⁻⁵,0.08]。")
    doc.add_heading("8.2 IMC/SIMC PI", level=2)
    add_formula(doc, "Kp=τ/[K(λ+L)]；Ti=min[τ,4(λ+L)]；Ki=Kp/Ti")
    add_para(doc, "λ是希望的闭环速度。λ大，控制更保守平滑；λ小，响应更快但对模型误差和延迟更敏感。项目不是拿默认保守IMC与AI比较，而是在训练工况中只搜索λ，再冻结公式得到的Kp、Ki。")
    add_table(doc, ["本次IMC量", "数值", "含义"], [
        ["λ", f(selected_imc["lambda_minutes"], 6) + " min", "训练集选出的闭环速度"],
        ["Kp", f(selected_imc["candidate_kp"], 7), "最终IMC比例增益，也是AI回退基准"],
        ["Ki", f(selected_imc["candidate_ki"], 7), "最终IMC积分增益"],
    ], [2000, 2100, 5260])
    add_callout(doc, "公平比较原则", "固定PI、FNN和RL必须面对同一对象、同一噪声、同一执行器限制和同一未见工况。基线也应合理调参，否则AI只是在击败一个故意设置很差的对手。")

    chapter(doc, "9. 贝叶斯优化BO：怎样少试几次找到Kp、Ki")
    add_para(doc, "BO适合‘一次仿真很贵、参数维度很少’的问题。这里输入x=[log Kp, log Ki]，输出是仿真综合目标J。使用对数空间是因为Ki跨越多个数量级，线性均匀抽样会浪费大量点。")
    add_numbers(doc, [
        "用IMC、Z-N和随机点得到初始真实仿真结果。",
        "高斯过程GP根据已知点估计每个候选的预测均值μ(x)和不确定度σ(x)。",
        "期望改进EI同时奖励预测值更好和不确定度更大，平衡利用与探索。",
        "选择EI最大的候选，运行真实3R2C仿真得到J，再更新GP。",
        "达到迭代预算后取真实评价中J最小的Kp、Ki；上线后参数固定，不在MCU运行GP。",
    ])
    add_formula(doc, "z=[Jbest-μ(x)-ξ]/σ(x)")
    add_formula(doc, "EI(x)=[Jbest-μ(x)-ξ]Φ(z)+σ(x)φ(z)")
    add_table(doc, ["本次全局BO输出", "数值"], [["Kp", f(bo["kp"], 7)], ["Ki", f(bo["ki"], 7)], ["训练目标", f(bo["objective"], 7)], ["真实批量评价次数", f(bo["evaluations"])]], [4200, 5160])
    add_callout(doc, "BO与FNN/RL的关系", "BO一方面产生一套全局固定PI；另一方面为每个训练工况产生较优Kp、Ki标签供FNN拟合。BO不是RL，也不等于在线自整定。")

    chapter(doc, "10. 数据集、训练集、验证集和测试集")
    add_para(doc, "本项目有数据集，但目前全部来自仿真，不是实测BMS数据。‘有CSV’不等于‘有真实数据’。最终实验使用48个分层训练工况和16个不同随机种子的独立留出工况。训练48个内部再切成38个拟合工况和10个内部验证工况。")
    add_table(doc, ["数据", "数量", "用途", "文件"], [
        ["训练工况", "48", "IMC/BO训练；FNN/RL再内部切分", "training_scenarios.csv"],
        ["FNN拟合工况", "38", "生成BO标签并拟合规则", "training_labels.csv"],
        ["内部验证工况", "10", "选FNN先验、RL早停和部署门", "训练历史CSV末行"],
        ["FNN逐状态样本", "2318", "e、ė到Kp/Ki标签", "fnn_training_samples.csv"],
        ["RL策略决策", "36000", "s,a,r,s′和TD误差", "rl_training_transitions.csv"],
        ["独立留出工况", "16", "最终泛化报告，不参与训练", "holdout_scenarios.csv"],
    ], [1800, 1000, 3200, 3360])
    doc.add_heading("10.1 分层工况为什么必要", level=2)
    add_bullets(doc, [
        "普通热启动：e>0且逐渐下降。",
        "冷启动：e<0，压缩机应停机，房间自然回暖。",
        "设定值下调：误差突然增大，ė为正。",
        "设定值上调：误差突然减小，ė为负，并考察卸载/停机。",
        "开门热脉冲：负荷突然增加后又撤销，覆盖两个变化率方向。",
        "持续负荷：考察慢扰动和容量不足。",
        "冷房间短时下调设定值：覆盖旧训练中缺失、但仍物理可解释的角落状态。",
    ])
    add_callout(doc, "仍然缺什么", "真实数据至少应包含时间戳、室温/送回风温度、设定值、请求与实际频率、室外温度、负荷代理、压力/排温、启停和故障码。还需要设备型号、房间体积和标定记录。当前仿真结果不能替代这些数据。", PALE_GOLD)

    chapter(doc, "11. FNN：它不是深度神经网络，而是可训练的模糊规则表")
    add_para(doc, "项目中的FNN是零阶TSK模糊规则面。输入e和ė各有5个中心，得到25条规则；输出不是压缩机命令，而是Kp、Ki。每次每个轴只有两个相邻隶属度非零，因此只计算4条规则。")
    add_formula(doc, "wᵢⱼ=μᵢ(e)·μⱼ(ė)，Σwᵢⱼ=1")
    add_formula(doc, "Kp=ΣwᵢⱼKpᵢⱼ；Ki=ΣwᵢⱼKiᵢⱼ")
    doc.add_heading("11.1 训练过程逐步展开", level=2)
    add_numbers(doc, [
        "对每个训练工况运行BO，得到该工况较优的固定Kp、Ki标签。",
        "用这组标签重新运行3R2C，每5 min采样e，并计算ė=Δe/Δt。",
        "找到e和ė相邻的两个中心，形成最多4个双线性权重。",
        "在log(Kp)、log(Ki)空间按权重累计，避免大数值标签支配算术平均。",
        "对每个规则加入IMC先验，未覆盖规则准确回到IMC。",
        "在10个内部验证工况比较多种先验权重1、2、5、10、25、50，选择验证目标最好者。",
        "只有规则覆盖≥80%且平均验证目标不劣于IMC才允许部署；否则保存候选表，但部署表写成IMC。",
    ])
    add_table(doc, ["最终FNN证据", "数值", "判读"], [
        ["覆盖", f"{int(float(fnn_final['occupied_rules']))}/25 = {f(fnn_final['rule_coverage_pct'])}%", "覆盖问题已解决"],
        ["拟合样本", f(fnn_final["state_samples"]), "38个工况每5 min采样"],
        ["验证IMC J", f(fnn_final["validation_baseline_objective"], 7), "安全基线"],
        ["最佳候选 J", f(fnn_final["validation_learned_objective"], 7), "高于IMC，未带来改善"],
        ["先验权重", f(fnn_final["selected_prior_weight"]), "验证选择的最保守候选"],
        ["部署结论", "拒绝，回退IMC", "覆盖充分不等于性能合格"],
    ], [1900, 2700, 4760])
    add_callout(doc, "为什么全覆盖仍失败", "不同热容量、热阻、制冷量和延迟的工况可能在同一(e,ė)位置需要不同增益。仅用两个输入存在不可消除的状态混叠；BO标签本身也有噪声。把规则填满只解决‘没见过’，没有解决‘同一个格子答案互相冲突’。", PALE_GOLD)

    chapter(doc, "12. RL：从状态、动作、奖励到Q-learning")
    add_para(doc, "强化学习不是拿一张固定标签表做监督学习，而是在虚拟环境里执行动作、观察后果并累计长期奖励。本项目使用表格Q-learning，不是深度RL，也不允许在真实压缩机上随机探索。")
    doc.add_heading("12.1 状态为什么必须尽量满足马尔可夫性", level=2)
    add_para(doc, "Q-learning假设当前状态包含决定未来的重要信息。旧版动作递归乘当前Kp/Ki，但状态只有e、Δe；同一个格子可能对应完全不同的当前增益，最优动作不唯一。最终修正把动作改成相对IMC的绝对目标，并加入上一次实际容量的停机/部分/高负荷3档。")
    add_formula(doc, "s=[bin(e), bin(ė), bin(u_applied)]")
    add_formula(doc, "Q表形状=5×5×3×9；MCU只部署75个动作索引，不部署675个浮点Q值")
    doc.add_heading("12.2 九个动作", level=2)
    add_para(doc, "Kp目标比例和Ki目标比例各取{0.75,1.0,1.3}，笛卡尔积得到9个动作。例如动作(1.3,0.75)表示目标Kp=1.3×IMC_Kp、目标Ki=0.75×IMC_Ki。实际应用仍每次最多变化±10%。")
    doc.add_heading("12.3 奖励的每一项", level=2)
    add_formula(doc, "r_interval=mean[-|e|-1.5·max(|e|-0.5,0)-0.08|Δu|-0.02u]")
    add_formula(doc, "r=r_interval-0.08·增益实际变化-0.015·目标比例偏离1")
    add_para(doc, "第一项要求温度接近目标；第二项额外惩罚超出±0.5℃；第三项抑制指令跳变；第四项是容量/能量代理；后两项阻止为了短期温度收益频繁或过度改参。奖励权重仍是工程选择。")
    doc.add_heading("12.4 Q更新公式从哪里来", level=2)
    add_formula(doc, "目标值 target = r + γ·maxₐ′Q(s′,a′)")
    add_formula(doc, "TD误差 δ = target-Q(s,a)")
    add_formula(doc, "Q(s,a)←Q(s,a)+α·δ")
    add_para(doc, "α=0.12表示本次新证据修正12%的TD差；γ=0.94表示未来奖励仍重要；ε从0.25降到0.03，早期多探索、后期多利用。回合最后一步没有未来项。")
    doc.add_heading("12.5 安全动作屏蔽", level=2)
    add_para(doc, "当房间已经明显低于设定值时，提高制冷PI增益没有帮助，只会让下一次启动更激进，因此这些动作在训练和部署中都被屏蔽。接近设定值时也屏蔽提高Ki的动作。屏蔽不是替代RL，而是把已知工程禁区从学习空间移除。")
    add_table(doc, ["最终RL证据", "数值", "判读"], [
        ["训练回合", f(rl_final["episode"]), "每回合最多48次五分钟决策"],
        ["对象步", f(rl_final["environment_steps"]), "180000个一分钟仿真步"],
        ["策略决策", "36000", "实际Q更新次数"],
        ["25热状态覆盖", f"{rl_final['visited_states']}/25 = {f(rl_final['state_coverage_pct'])}%", "覆盖问题已解决"],
        ["75完整组合自然覆盖", f"{rl_final['visited_full_states']}/75 = {f(rl_final['full_state_coverage_pct'])}%", "未伪造不安全组合"],
        ["验证IMC J", f(rl_final["validation_baseline_objective"], 7), "安全基线"],
        ["RL验证 J", f(rl_final["validation_learned_objective"], 7), "优于门槛，内部通过"],
        ["部署结论", "通过", "导出冻结策略，上线ε=0"],
    ], [2100, 2900, 4360])

    chapter(doc, "13. 为什么旧版只覆盖FNN 48%、RL 16/25")
    add_figure(doc, coverage_fig, "图3  整改前后状态/规则覆盖对比")
    add_table(doc, ["旧问题", "直接后果", "最终修正"], [
        ["初温只等于或高于设定值", "负误差状态很少", "加入冷启动、设定值上下调"],
        ["FNN每分钟Δe却用±1℃宽中心", "样本挤在中间变化率列", "统一为ė=Δe/Δt并重设中心"],
        ["RL训练用5 min Δe，运行用最近1 min Δe", "训练/部署状态含义不同", "外环只在更新时保存误差并计算℃/min"],
        ["RL递归乘当前增益但状态不含当前增益", "同一状态对应不同未来", "绝对IMC目标，消除隐藏增益"],
        ["状态不含执行器模式", "停机和高负荷被混在一起", "增加3档上一次实际容量"],
        ["只保存汇总历史", "无法审计每个样本/转移", "保存FNN逐状态样本和RL 36000条转移"],
    ], [2350, 3000, 4010])
    add_callout(doc, "覆盖率的正确理解", "覆盖率是数据支持证据，不是性能证据。25/25说明每个热状态至少到达过；它不说明每个动作都充分探索、不说明Q值收敛，也不说明留出工况优于IMC。")

    chapter(doc, "14. 验收为什么以前失败，最终怎样处理")
    add_para(doc, "旧RL在1个内部验证工况上J=56.88，而IMC为51.02，差约11.5%，所以正确地被拒绝。主要原因是训练工况单一、状态定义不一致、递归动作破坏状态完整性、覆盖不足以及没有早停。")
    add_numbers(doc, [
        "先保证数据是物理轨迹，不用人为伪状态凑覆盖。",
        "把训练课程扩到冷热启动、设定值变化和负荷扰动。",
        "把Δe统一为误差变化率ė，消除采样周期依赖。",
        "把RL动作改为绝对IMC目标，并加入实际容量模式。",
        "每隔约25回合在10个内部验证工况回放，保存最佳Q表，而不是盲目使用最后一回合。",
        "覆盖≥80%且验证J不超过IMC的102%才允许RL部署。",
        "FNN门槛更严格：覆盖≥80%且验证J不得劣于IMC；否则候选和部署产物分开。",
    ])
    add_callout(doc, "为什么不直接降低门槛", "降低门槛只能把失败改名为通过，不能改善控制。最终FNN仍未通过，因此部署表就是IMC；这是安全功能，不是算法失败被隐藏。", PALE_RED)

    chapter(doc, "15. 48训练+16留出工况的最终结果")
    comp = []
    for name in ["Ziegler-Nichols", "IMC PI", "Bayesian Auto-tune", "FNN Self-tuning PI", "RL Self-tuning PI"]:
        item = summary_by_name[name]
        comp.append([
            name,
            f(item["mean_objective"], 6),
            f(item["median_objective"], 6),
            f(item["mean_iae_c_hour"], 5),
            f(item["mean_itae_c_hour2"], 5),
            f(item["mean_settling_time_hour"], 4),
            f(item["stable_rate"], 3),
        ])
    add_table(doc, ["控制器", "平均J", "中位J", "平均IAE", "平均ITAE", "调节h", "稳定率"], comp, [2200, 1120, 1120, 1220, 1220, 1200, 1280])
    imc_mean = float(summary_by_name["IMC PI"]["mean_objective"])
    rl_mean = float(summary_by_name["RL Self-tuning PI"]["mean_objective"])
    add_callout(doc, "数字应该怎样说", f"FNN实际部署回退IMC，所以结果与IMC完全相同。RL平均J={rl_mean:.4f}，IMC={imc_mean:.4f}，RL仅改善{100*(imc_mean-rl_mean)/imc_mean:.2f}%；但RL平均IAE、ITAE和调节时间并非全面更好。正确结论是‘通过当前综合门槛、优势很小且需要真实数据复核’，不是‘RL明显优于PI’。", PALE_GOLD)
    for picture, caption in (
        (OUT / "fnn_training_trace.png", "图4  FNN覆盖、拟合误差和规则参数轨迹"),
        (OUT / "rl_training_trace.png", "图5  RL奖励、TD误差、覆盖率和增益轨迹"),
        (OUT / "training_convergence_overview.png", "图6  BO、FNN和RL不同含义的收敛证据"),
        (OUT / "dynamic_comparison.png", "图7  混合动态工况下五种控制器波形"),
    ):
        if picture.exists():
            add_figure(doc, picture, caption)

    chapter(doc, "16. 从Python训练结果到MCU：输入、输出和时序")
    add_figure(doc, control_flow, "图8  嵌入式运行时闭环和安全顺序")
    add_table(doc, ["任务", "周期", "输入", "输出", "为什么分频"], [
        ["传感器/PI", "100 ms", "温度、设定值、dt", "容量请求", "快速控制和准时性"],
        ["FNN/RL外环", "示例2 s", "e、ė、上次实际容量", "候选Kp/Ki", "热系统慢，不需每100 ms改参"],
        ["执行器限制", "每个PI周期", "请求、状态、计时", "可执行容量", "所有命令都必须安全"],
        ["遥测", "0.5～2 s", "温度、命令、增益、故障", "串口/总线日志", "调试和追溯"],
    ], [1600, 1200, 2700, 2100, 1760])
    doc.add_heading("16.1 自动导出解决什么", level=2)
    add_para(doc, "旧C++硬编码的是早期FNN/RL表、错误覆盖掩码、±25%增益变化和递归增量动作，已经与Python分叉。新增export_policy.py从最终目录读取验收标志和部署表，生成generated_policy.hpp：")
    add_bullets(doc, [
        "FNN/RL是否通过验收。",
        "IMC回退Kp、Ki。",
        "FNN中心、规则后件。",
        "RL误差/变化率/容量边界、9个绝对目标比例、75个动作索引和覆盖掩码。",
        "产物版本和CRC32，便于固件拒绝损坏或错版本参数。",
    ])
    add_formula(doc, "ė=[e(k)-e(k-1)]/Δt，单位°C/min", "这样5 min训练和2 s嵌入式周期使用同一物理量。")
    add_callout(doc, "本次导出状态", "FNN accepted=false，C++收到的是IMC回退表；RL accepted=true，导出冻结策略；PC SIL通过。CRC32=0x76E995D6。")

    chapter(doc, "17. SIL、开发板、HIL和实机分别证明什么")
    add_table(doc, ["层级", "连接方式", "能证明", "不能证明", "本项目状态"], [
        ["Python仿真", "全软件", "算法逻辑和批量比较", "真实设备性能", "完成"],
        ["PC SIL", "C++控制器+虚拟对象", "查表、限幅、调度、异常回退", "目标MCU资源/WCET", "通过"],
        ["Wokwi", "虚拟STM32/ESP32", "目标编译、基础外设入口", "真实模拟精度和时序", "部分"],
        ["真实开发板台架", "MCU+信号源/负载", "ADC/PWM/通信、Flash/RAM/WCET", "真实热闭环", "未完成"],
        ["HIL", "真实MCU+实时对象仿真器", "闭环、故障注入、掉电/超时", "整机认证", "未开始"],
        ["实机", "真实传感器/变频器/压缩机", "真实性能和安全协同", "跨型号泛化", "未开始"],
    ], [1250, 2200, 2200, 2200, 1510])
    add_callout(doc, "PC纳秒数不能当MCU WCET", "PC SIL记录的最坏PI/AI纳秒只用于发现明显算法异常。STM32/ESP32必须用目标编译器、真实时钟、DWT/周期计数器或GPIO+逻辑分析仪重新测量，并记录中断抖动和最大栈。", PALE_GOLD)

    chapter(doc, "18. 真实BMS数据接入和安全调试流程")
    add_table(doc, ["字段", "最低要求", "用途"], [
        ["timestamp", "单调时间戳和采样周期", "计算ė、对齐事件"],
        ["Tzone/Treturn/Tsupply", "单位、位置、标定信息", "温度状态与传感器诊断"],
        ["setpoint", "实际生效设定值", "计算误差"],
        ["u_request/u_applied", "请求与实际频率分开", "识别执行器限制/延迟"],
        ["Tout/load proxy", "室外温度、功耗/占用", "区分扰动"],
        ["pressure/discharge temp", "量程、报警阈值", "独立安全联锁"],
        ["mode/alarms", "启停、除霜、故障码", "筛除非正常数据"],
    ], [2100, 3750, 3510])
    add_numbers(doc, [
        "先做数据质量检查：缺失、重复、时间跳变、传感器卡死、单位错误。",
        "只选择稳定模式或明确事件窗口；排除维修、除霜和报警期间的数据。",
        "用安全批准的小阶跃或自然运行数据拟合FOPDT/3R2C，并报告置信区间与残差。",
        "先运行固定IMC，在真实设备建立基线和联锁记录。",
        "AI先影子运行：计算建议但不下发，比较其建议与工程师/基线。",
        "有限权限试验：限制最大增益偏差、容量范围和试验时段，有人工急停。",
        "故障注入：传感器断线、卡值、通信超时、压力报警、看门狗复位和参数CRC错误。",
        "只有多天/多负荷/多噪声重复通过后，才讨论扩大权限；保留远程回退和版本回滚。",
    ])

    chapter(doc, "19. 参数字典：每个数值从哪里来")
    add_table(doc, ["参数", "最终值/范围", "来源", "意义"], [
        ["dt模型", "1 min", "批量仿真设置", "数值积分步长，不是MCU周期"],
        ["PI周期", "100 ms", "嵌入式设计目标", "快速内环调度"],
        ["AI周期", "Python 5 min；MCU示例2 s", "训练效率/部署设计", "通过ė归一化保持语义"],
        ["Kp边界", "0.002～1.5", "工程安全边界", "限制比例作用"],
        ["Ki边界", "1e-5～0.08", "工程安全边界", "限制积分作用"],
        ["增益单次变化", "±10%", "安全层", "防止在线突变"],
        ["FNN中心e", "-3,-0.75,0,1.5,5℃", "分层轨迹范围", "覆盖冷热和偏热"],
        ["FNN中心ė", "-0.30,-0.05,0,0.05,0.30℃/min", "分层轨迹范围", "统一不同节拍"],
        ["RL α", "0.12", "训练超参数", "TD更新步长"],
        ["RL γ", "0.94", "训练超参数", "未来奖励折扣"],
        ["RL ε", "0.25→0.03", "训练探索计划", "上线固定为0"],
        ["RL动作比例", "0.75/1/1.3", "安全目标集合", "相对IMC的绝对目标"],
        ["FNN门槛", "覆盖≥80%，J≤IMC", "部署验收", "不允许性能退化"],
        ["RL门槛", "覆盖≥80%，J≤1.02IMC", "部署验收", "允许2%噪声/近似容差"],
    ], [1800, 2200, 2300, 3060])

    chapter(doc, "20. 常见误解与准确回答")
    add_table(doc, ["误解", "准确回答"], [
        ["覆盖100%就训练好了", "错误。覆盖只说明到达过；还要看每格样本数、动作覆盖、TD误差、验证和独立测试。"],
        ["FNN就是神经网络", "本项目是零阶TSK模糊规则表，没有多层网络和反向传播。"],
        ["RL直接控制压缩机", "错误。RL只给Kp/Ki目标，PI和独立安全层决定最终容量。"],
        ["BO在MCU开机后不断调参", "错误。BO离线运行，MCU通常只保存结果。"],
        ["FNN全覆盖所以应部署", "错误。最终候选验证不优于IMC，因此自动回退。"],
        ["RL平均J略好就证明AI优越", "错误。改善仅约0.37%，其他指标并非全面更好，且仍无真实数据。"],
        ["168 h说明空调要等7天", "这是电脑加速模拟的虚拟辨识时长；真实调试应使用历史BMS或安全短试验。"],
        ["4 h到目标一定不合理", "需要结合热负荷、制冷量、热容量、设定温差和执行器限制判断；模型参数未实测前不能当真实产品结论。"],
        ["PC SIL通过就能接压缩机", "错误。还缺目标板资源/WCET、HIL、联锁和实机安全验证。"],
    ], [3300, 6060])

    chapter(doc, "21. 复现命令、产物和验收清单")
    add_formula(doc, 'python main.py --train-samples 48 --test-samples 16 --bo-iterations 4 --seed 23 --output outputs_adaptive_final_v2')
    add_formula(doc, 'python embedded/export_policy.py outputs_adaptive_final_v2')
    add_formula(doc, 'python embedded/wokwi/prepare_projects.py')
    add_formula(doc, 'python embedded/run_mcu_validation.py')
    add_formula(doc, 'python -m unittest discover -s tests -v')
    add_table(doc, ["验收项", "当前结果", "下一步"], [
        ["单元/流水线测试", "10/10通过", "加入Python-C++黄金向量逐点比对"],
        ["FNN覆盖", "25/25", "用真实数据重训；候选当前拒绝"],
        ["RL热状态覆盖", "25/25", "扩大真实/高保真验证"],
        ["独立留出", "16工况；RL仅小幅改善", "多随机种子、置信区间、不同设备"],
        ["自动导出", "版本+CRC+验收状态", "在MCU启动时实际校验CRC"],
        ["PC SIL", "通过", "增加限制器黄金测试"],
        ["目标板ROM/RAM/WCET", "未完成", "STM32/ESP32真实编译和测量"],
        ["HIL/实机", "未开始", "联锁和故障注入后再上机"],
    ], [2800, 2600, 3960])

    chapter(doc, "22. 最终工程结论", page_break_before=True)
    add_bullets(doc, [
        "旧FNN 48%和RL 16/25覆盖的根因不是随机运气，而是数据分布、变化率尺度和状态定义错误。",
        "分层物理工况和统一ė后，FNN与RL热状态均达到25/25；没有使用伪状态。",
        "FNN候选虽然全覆盖，但只用e、ė无法充分区分不同对象，内部验证不优于IMC；工程上正确动作是拒绝部署。",
        "RL改成绝对IMC目标并加入实际容量模式，修复主要状态混叠；内部通过，16留出平均J仅改善约0.37%。",
        "自动导出、版本/CRC、验收标志、±10%增益限制和压缩机限制器已同步到C++，PC SIL通过。",
        "本项目仍是仿真和PC SIL成果，不是已经验证的真实空调控制产品。下一关键路径是真实BMS数据、目标板资源/WCET、HIL和受控实机试验。",
    ])
    add_callout(doc, "最稳妥的当前部署选择", "真实硬件首轮应使用通过调参的固定IMC PI和完整独立联锁。RL只能作为受限、可即时回退的候选；FNN候选暂不部署。", GREEN)

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build())
