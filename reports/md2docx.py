# -*- coding: utf-8 -*-
"""通用 Markdown → docx 转换器（中文工程报告排版）。

用法：
    python reports/md2docx.py                          # 默认转换 实验报告生成.md
    python reports/md2docx.py 实验报告_AI-PI整定.md     # 转换指定文件
    python reports/md2docx.py src.md out.docx          # 指定输出

支持：#~#### 标题、表格（Table Grid + 表头底纹）、无序/有序列表（含缩进子项）、
引用块、![图片](路径)、**加粗**、`行内代码`、```代码块```、--- 分隔线，
以及 LaTeX 公式：$$...$$ 独立公式块与 $...$ 行内公式（经 matplotlib mathtext
渲染为 PNG 图片嵌入，缓存于 reports/figures/equations/；渲染失败时回退为
Cambria Math 斜体文本并打印警告）。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

EAST_ASIA = "微软雅黑"
CODE_FONT = "Consolas"
MATH_FONT = "Cambria Math"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x59, 0x59, 0x59)
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\$[^$\n]+\$)")

MATH_CACHE = Path(__file__).resolve().parent / "figures" / "equations"
MATH_DPI = 300


def render_math_png(latex: str, path: Path, fontsize: float) -> bool:
    """用 matplotlib mathtext 把 LaTeX 公式渲染为紧致裁剪的透明底 PNG。"""
    if any("一" <= ch <= "鿿" for ch in latex):
        print(f"[warn] 公式含中文字符，mathtext 无法渲染: {latex[:50]}")
        return False
    import matplotlib
    matplotlib.use("Agg")
    matplotlib.rcParams["mathtext.fontset"] = "cm"
    import matplotlib.pyplot as plt
    path.parent.mkdir(parents=True, exist_ok=True)
    fig = None
    try:
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, "$" + latex + "$", fontsize=fontsize, color="#1a1a1a")
        fig.savefig(str(path), dpi=MATH_DPI, bbox_inches="tight",
                    pad_inches=0.02, transparent=True)
        return True
    except Exception as exc:
        print(f"[warn] 公式渲染失败: {latex[:60]} ({exc})")
        return False
    finally:
        if fig is not None:
            plt.close(fig)


def math_png_size_cm(path: Path, scale: float = 1.0):
    """按 300 dpi 换算图片物理尺寸（保持公式字号与正文一致）。"""
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    return w / MATH_DPI * 2.54 * scale, h / MATH_DPI * 2.54 * scale


import itertools

_eq_counter = itertools.count(1)


def add_display_math(doc, latex: str) -> None:
    """独立公式块：渲染为 PNG 居中嵌入；失败回退为 Cambria Math 斜体文本。"""
    path = MATH_CACHE / f"eq_{next(_eq_counter):03d}.png"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if render_math_png(latex, path, fontsize=13):
        _, h = math_png_size_cm(path)
        p.add_run().add_picture(str(path), height=Cm(h))
    else:
        run = p.add_run(latex)
        set_run_font(run, size=11, italic=True, font=MATH_FONT, ascii_font=MATH_FONT)


def split_row(row):
    """按未转义的 | 分割单元格，并把 \\| 还原为 |。"""
    parts = re.split(r"(?<!\\)\|", row.strip().strip("|"))
    return [p.strip().replace("\\|", "|") for p in parts]


def set_run_font(run, size=10.5, bold=False, italic=False, color=None,
                 font=EAST_ASIA, ascii_font=None):
    run.font.name = ascii_font or font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def add_inline(par, text, size=10.5, base_bold=False, color=None, _depth=0):
    for seg in INLINE_RE.split(text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            if _depth < 4:
                add_inline(par, seg[2:-2], size=size, base_bold=True,
                           color=color, _depth=_depth + 1)
            else:
                run = par.add_run(seg[2:-2])
                set_run_font(run, size=size, bold=True, color=color)
        elif seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
            run = par.add_run(seg[1:-1])
            set_run_font(run, size=max(size - 1, 8), font=EAST_ASIA,
                         ascii_font=CODE_FONT, color=RGBColor(0xA3, 0x15, 0x15))
        elif seg.startswith("$") and seg.endswith("$") and len(seg) > 2:
            latex = seg[1:-1]
            path = MATH_CACHE / f"eq_{next(_eq_counter):03d}.png"
            if render_math_png(latex, path, fontsize=11):
                _, h = math_png_size_cm(path)
                par.add_run().add_picture(str(path), height=Cm(h))
            else:
                run = par.add_run(latex)
                set_run_font(run, size=size, italic=True, font=MATH_FONT,
                             ascii_font=MATH_FONT, color=color)
        else:
            run = par.add_run(seg)
            set_run_font(run, size=size, bold=base_bold, color=color)


def add_table(doc, header, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    for j, cell_text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(cell.paragraphs[0], cell_text, size=9, base_bold=True)
        shd = cell._element.get_or_add_tcPr().makeelement(
            qn("w:shd"), {qn("w:fill"): "DEEAF6", qn("w:val"): "clear"})
        cell._element.get_or_add_tcPr().append(shd)
    for i, row in enumerate(rows):
        row = (row + [""] * len(header))[:len(header)]
        for j, cell_text in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            add_inline(cell.paragraphs[0], cell_text, size=9)
    doc.add_paragraph()


def convert(src: Path, dst: Path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    lines = src.read_text(encoding="utf-8").splitlines()
    i, n = 0, len(lines)
    in_code = False
    code_buf = []

    def flush_code():
        nonlocal code_buf
        if code_buf:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            for k, ln in enumerate(code_buf):
                if k:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.5)
                run = p.add_run(ln if ln else " ")
                set_run_font(run, size=9, font=EAST_ASIA, ascii_font=CODE_FONT)
            code_buf = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                flush_code()
            in_code = not in_code
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 独立公式块 $$...$$
        if stripped.startswith("$$"):
            if stripped.endswith("$$") and len(stripped) > 4:
                add_display_math(doc, stripped[2:-2].strip())
                i += 1
                continue
            buf = []
            i += 1
            while i < n and not lines[i].strip().endswith("$$"):
                buf.append(lines[i].strip())
                i += 1
            add_display_math(doc, " ".join(buf))
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < n and re.match(
                r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = split_row(stripped)
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, header, rows)
            continue

        if not stripped or stripped == "---":
            i += 1
            continue

        # 图片
        m = re.match(r"^!\[(.*)\]\((.+)\)$", stripped)
        if m:
            img = src.parent / m.group(2)
            if img.exists():
                doc.add_picture(str(img), width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(cap, m.group(1), size=9, color=GRAY)
            else:
                p = doc.add_paragraph()
                add_inline(p, f"[缺图: {m.group(2)}]", size=9, color=GRAY)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, text, size=20, base_bold=True, color=ACCENT)
                doc.add_paragraph()
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                add_inline(p, text, size={2: 15, 3: 12.5, 4: 11}.get(level, 11),
                           base_bold=True, color=ACCENT)
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            quote = stripped.lstrip("> ").rstrip()
            if quote:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, quote, size=9.5, color=GRAY)
            i += 1
            continue

        # 无序列表
        m = re.match(r"^(\s*)-\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75 + (0.6 if indent >= 2 else 0))
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run("•  " if indent < 2 else "–  ")
            set_run_font(run, size=10.5, color=ACCENT)
            add_inline(p, m.group(2))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75 + (0.6 if indent >= 2 else 0))
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(f"{m.group(2)}. ")
            set_run_font(run, size=10.5, bold=True, color=ACCENT)
            add_inline(p, m.group(3))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, stripped)
        i += 1

    flush_code()
    dst.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(dst))
    except PermissionError:
        dst = dst.with_suffix(".new.docx")
        doc.save(str(dst))
        print(f"[warn] 目标文件被占用（可能已在 Word 中打开），已另存为: {dst.name}")
    return dst


def main():
    args = sys.argv[1:]
    here = Path(__file__).resolve().parent
    if not args:
        src, dst = here / "实验报告生成.md", here / "实验报告生成.docx"
    elif len(args) == 1:
        src = here / args[0] if not Path(args[0]).is_absolute() else Path(args[0])
        dst = src.with_suffix(".docx")
    else:
        src = Path(args[0]) if Path(args[0]).is_absolute() else here / args[0]
        dst = Path(args[1]) if Path(args[1]).is_absolute() else here / args[1]
    out = convert(src, dst)
    sys.stdout.buffer.write(f"已生成: {out}\n".encode("utf-8"))


if __name__ == "__main__":
    main()
